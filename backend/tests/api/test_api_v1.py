"""API tests against an isolated temporary database (no dependency on data/seo.db)."""
import json
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text
from openpyxl import Workbook

from seo_brain.api import deps
from seo_brain.api.main import create_app
from seo_brain.api.routers import graph as graph_router
from seo_brain.api.routers import sites as sites_router
from seo_brain.database.db import connect as legacy_connect
from seo_brain.db.engine import make_engine
from seo_brain.db.migrate import migrate
from seo_brain.graph.model import GraphEdge, GraphNode
from seo_brain.graph.store import SqlGraphStore


@pytest.fixture
def client(tmp_path, monkeypatch):
    dbfile = tmp_path / "api.db"
    eng = make_engine("sqlite:///" + dbfile.as_posix())
    migrate(eng)
    monkeypatch.delenv("API_TOKEN", raising=False)
    monkeypatch.setattr(graph_router, "connect", lambda: legacy_connect(dbfile))   # legacy analytics → same temp DB
    monkeypatch.setattr(sites_router, "PROJECT_ROOT", tmp_path)                       # workspaces under tmp, not data/
    app = create_app()
    app.dependency_overrides[deps.engine] = lambda: eng
    from seo_brain.ai.gateway import Gateway as _GW
    app.dependency_overrides[deps.gateway] = (lambda g: (lambda: g))(_GW(eng))   # isolate: never the live DB gateway
    c = TestClient(app)
    c.eng = eng  # type: ignore[attr-defined]
    return c


def _seed(c: TestClient):
    r = c.post("/api/v1/sites", json={"site_id": "demo", "name": "Demo", "canonical_url": "https://demo.example/",
                                       "wp_url": "https://demo.example", "business_type": "auto-service"})
    assert r.status_code == 201, r.text
    store = SqlGraphStore(c.eng)
    store.upsert_nodes([GraphNode("site:demo", "demo", "SITE", {"label": "Demo"}),
                        GraphNode("page:https://demo.example/a", "demo", "PAGE", {"label": "A", "url": "https://demo.example/a"}),
                        GraphNode("query:امداد", "demo", "QUERY", {"label": "امداد"})])
    store.upsert_edges([GraphEdge("site:demo", "page:https://demo.example/a", "HAS_PAGE", site_id="demo"),
                        GraphEdge("page:https://demo.example/a", "query:امداد", "RANKS_FOR", 0.5, {"props": {"position": 8}}, "demo")])


def test_health_reports_migrations(client):
    r = client.get("/api/v1/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok" and "0002" in body["migrations"]["applied"] and body["migrations"]["pending"] == []


def test_call_attribution_uses_actual_call_time_and_preserves_manual_choice(client):
    _seed(client)
    at = datetime.now(timezone.utc).replace(microsecond=0) - timedelta(days=1)
    ts = at.isoformat()
    with client.eng.begin() as cx:
        cx.execute(text("""INSERT INTO track_sessions(session_id,site_id,visitor_id,day,started_at,last_seen_at,channel)
            VALUES ('organic-session','demo','visitor',:day,:ts,:ts,'organic')"""), {"day": ts[:10], "ts": ts})
        cx.execute(text("""INSERT INTO track_events(site_id,session_id,day,ts,type,label)
            VALUES ('demo','organic-session',:day,:ts,'tel_click','02100000000')"""), {"day": ts[:10], "ts": ts})
    created = client.post('/api/v1/call-center/calls', json={"site_id": "demo", "occurred_at": ts,
        "customer_name": "Test Caller", "phone": "09120000000"})
    assert created.status_code == 201, created.text
    row = created.json()
    assert row["source"] == "seo" and row["source_confidence"] == "probable" and row["auto_attributed"] == 1
    assert row["attribution_event"].startswith("track:")
    second = client.post('/api/v1/call-center/calls', json={"site_id": "demo", "occurred_at": ts,
        "customer_name": "Another Caller", "phone": "09120000002"})
    assert second.status_code == 201 and second.json()["source"] == "unknown"
    # A second possible caller click makes the identity ambiguous on recheck.
    with client.eng.begin() as cx:
        cx.execute(text("""INSERT INTO track_events(site_id,session_id,day,ts,type,label)
            VALUES ('demo','organic-session',:day,:ts,'tel_click','02100000000')"""), {"day": ts[:10], "ts": ts})
    reconciled = client.post('/api/v1/call-center/reconcile', params={"site_id": "demo"})
    assert reconciled.status_code == 200 and reconciled.json()["changed"] == 1
    row = client.get('/api/v1/call-center/calls', params={"site_id": "demo"}).json()["items"][0]
    assert row["source"] == "unknown" and row["auto_attributed"] == 0
    manual = client.patch(f'/api/v1/call-center/calls/{row["id"]}', json={"source": "ads", "source_basis": "manual"})
    assert manual.status_code == 200 and manual.json()["auto_attributed"] == 0
    client.post('/api/v1/call-center/reconcile', params={"site_id": "demo"})
    assert client.get('/api/v1/call-center/calls', params={"site_id": "demo"}).json()["items"][0]["source"] == "ads"


def test_phone_required_for_manual_calls_and_edits(client):
    path = '/api/v1/call-center/calls'
    assert client.post(path, json={"customer_name": "Test Caller"}).status_code == 422
    assert client.post(path, json={"customer_name": "Test Caller", "phone": " + "}).status_code == 422
    created = client.post(path, json={"customer_name": "Test Caller", "phone": "۰۹۱۲۳۴۵۶۷۸۹"})
    assert created.status_code == 201 and created.json()["phone"] == "09123456789"
    assert client.patch(f'{path}/{created.json()["id"]}', json={"phone": ""}).status_code == 422


def test_late_click_reconciliation_marks_ads_probable(client):
    _seed(client)
    at = (datetime.now(timezone.utc) - timedelta(hours=4)).replace(microsecond=0).isoformat()
    created = client.post('/api/v1/call-center/calls', json={"site_id": "demo", "occurred_at": at,
        "phone": "09120000001"})
    assert created.status_code == 201 and created.json()["source"] == "unknown"
    with client.eng.begin() as cx:
        cx.execute(text("""INSERT INTO ads_click_events(event_uuid,site_id,event_type,received_at,ip_address,ip_hash,gclid)
            VALUES ('test-late-click','demo','tel_click',:at,'127.0.0.1','test-hash','test-gclid')"""), {"at": at})
    result = client.post('/api/v1/call-center/reconcile', params={"site_id": "demo"})
    assert result.status_code == 200 and result.json()["changed"] == 1
    row = client.get('/api/v1/call-center/calls', params={"site_id": "demo"}).json()["items"][0]
    assert row["source"] == "ads" and row["source_confidence"] == "probable"


def test_sites_crud_and_workspace(client):
    assert client.get("/api/v1/sites").json() == []
    _seed(client)
    sites = client.get("/api/v1/sites").json()
    assert [s["site_id"] for s in sites] == ["demo"] and sites[0]["mode"] == "manual"
    assert sites[0]["workspace_path"] == "data/sites/demo"
    r = client.patch("/api/v1/sites/demo", json={"mode": "assisted", "country": "IR"})
    assert r.status_code == 200 and r.json()["mode"] == "assisted" and r.json()["country"] == "IR"
    assert client.post("/api/v1/sites", json={"site_id": "demo", "name": "x", "canonical_url": "https://x.example/"}).status_code == 409
    assert client.get("/api/v1/sites/nope").status_code == 404
    assert client.post("/api/v1/sites", json={"site_id": "Bad Slug", "name": "x", "canonical_url": "https://x.example/"}).status_code == 422


def test_graph_endpoints_neo4j_shape(client):
    _seed(client)
    s = client.get("/api/v1/sites/demo/graph/summary").json()
    assert s["nodes"] == 3 and s["edges"] == 2 and s["by_relation_type"]["RANKS_FOR"] == 1
    n = client.get("/api/v1/sites/demo/graph/node/page:https://demo.example/a").json()
    assert set(n) == {"id", "site_id", "type", "metadata"} and n["type"] == "PAGE"
    sg = client.get("/api/v1/sites/demo/graph/subgraph", params={"center": "site:demo", "hops": 2}).json()
    assert len(sg["nodes"]) == 3 and len(sg["edges"]) == 2
    e = sg["edges"][0]
    assert set(e) == {"source", "target", "relation_type", "weight", "metadata", "site_id"}
    nb = client.get("/api/v1/sites/demo/graph/neighbors/query:امداد", params={"direction": "in"}).json()
    assert len(nb["edges"]) == 1
    assert client.get("/api/v1/sites/demo/graph/nodes", params={"types": "query"}).json()[0]["id"] == "query:امداد"
    sr = client.get("/api/v1/sites/demo/graph/search", params={"q": "امداد"}).json()
    assert [x["id"] for x in sr["nodes"]] == ["query:امداد"] and "fts" in sr
    assert client.get("/api/v1/sites/demo/graph/node/nope").status_code == 404
    assert client.get("/api/v1/sites/demo/graph/orphans").status_code == 200


def test_portfolio_overview_is_one_consistent_snapshot(client):
    _seed(client)
    body = client.get("/api/v1/portfolio/overview").json()
    assert body["totals"]["sites"] == 1
    assert body["totals"]["graph_nodes"] == 3 and body["totals"]["graph_edges"] == 2
    assert body["by_node_type"] == {"PAGE": 1, "QUERY": 1, "SITE": 1}
    assert body["sites"][0]["site_id"] == "demo"
    assert body["sites"][0]["state"] == "partial"
    assert body["sites"][0]["setup_progress"] == 50

    with client.eng.begin() as cx:
        cx.execute(text("INSERT INTO posts(site_id,wp_id,type,url,title,status) VALUES('demo',1,'post','https://demo.example/post','Post','publish')"))
        notes = json.dumps({"status": "succeeded", "progress": 1, "finished_at": "2026-08-20T10:00:00Z", "errors": []})
        cx.execute(text("INSERT INTO sync_runs(run_id,site_id,source,started_at,finished_at,status,notes) VALUES('run-1','demo','wordpress_pipeline','2026-08-20T09:00:00Z','2026-08-20T10:00:00Z','succeeded',:notes)"), {"notes": notes})

    body = client.get("/api/v1/portfolio/overview").json()
    assert body["totals"]["content"] == 1 and body["totals"]["ready_sites"] == 1
    assert body["sites"][0]["state"] == "ready"
    assert body["sites"][0]["state_reason"] and body["sites"][0]["next_action"]
    assert body["sites"][0]["setup_steps"]["graph_ready"] is True
    assert body["recent_activity"][0]["site_name"] == "Demo"

    # A stale local snapshot must not hide a broken WordPress connection.
    with client.eng.begin() as cx:
        cx.execute(text("INSERT INTO site_connections(site_id,kind,status,detail,tested_at) VALUES('demo','wordpress','error','{}','2026-08-20T11:00:00Z')"))
    body = client.get("/api/v1/portfolio/overview").json()
    assert body["sites"][0]["state"] == "attention"
    assert body["sites"][0]["next_action"] == "اصلاح اتصال وردپرس"
    assert body["sites"][0]["issues"][0]["severity"] == "blocking"


def test_portfolio_exposes_source_coverage_and_work_backlog(client):
    _seed(client)
    today = datetime.now(timezone.utc).date().isoformat()
    past = (datetime.now(timezone.utc) - timedelta(days=2)).isoformat()
    with client.eng.begin() as cx:
        cx.execute(text("""INSERT INTO gsc_property_daily
            (site_id,date,search_type,property,clicks,impressions,position,sync_run_id,observed_at)
            VALUES ('demo',:d,'web','sc-domain:demo.example',2,20,4,'g1',:t)"""), {"d": today, "t": past})
        cx.execute(text("""INSERT INTO work_items
            (site_id,title,status,due_at,created_at,updated_at)
            VALUES ('demo','Fix title','new',:due,:now,:now)"""), {"due": past, "now": past})
    body = client.get("/api/v1/portfolio/overview").json()
    assert body["sites"][0]["data_coverage"]["gsc"] == {"last_date": today, "days_28": 1}
    assert body["sites"][0]["data_coverage"]["ga4"] == {"last_date": None, "days_28": 0}
    assert body["sites"][0]["work"] == {"open": 1, "overdue": 1, "unassigned": 1}
    assert body["totals"]["overdue_work"] == 1


def test_call_center_csv_import_preview_mapping_and_deduplication(client):
    _seed(client)
    content = "مشتری,موبایل,کانال,شهر\nعلی,۰۹۱۲۳۴۵۶۷۸۹,ادز,تهران\nمینا,۰۹۳۵۱۲۳۴۵۶۷,گوگل,شیراز\n".encode("utf-8-sig")
    path = "/api/v1/call-center/calls/import"
    mapping = json.dumps({"customer_name": "مشتری", "phone": "موبایل", "source": "کانال", "region": "شهر"})
    upload = lambda: {"file": ("calls.csv", content, "text/csv")}
    preview = client.post(path, files=upload(), data={"mapping": mapping, "dry_run": "true", "default_site_id": "demo"})
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows_valid"] == 2 and preview.json()["rows_imported"] == 0
    assert [row["source"] for row in preview.json()["preview"]] == ["ads", "seo"]
    assert client.get("/api/v1/call-center/calls").json()["total"] == 0
    done = client.post(path, files=upload(), data={"mapping": mapping, "dry_run": "false", "default_site_id": "demo"})
    assert done.status_code == 200 and done.json()["rows_imported"] == 2
    repeat = client.post(path, files=upload(), data={"mapping": mapping, "dry_run": "false", "default_site_id": "demo"})
    assert repeat.json()["rows_skipped"] == 2 and repeat.json()["rows_imported"] == 0
    calls = client.get("/api/v1/call-center/calls").json()
    assert calls["total"] == 2 and {row["source_basis"] for row in calls["items"]} == {"import"}
    missing_phone = "مشتری,موبایل\nبدون شماره,\n".encode("utf-8-sig")
    rejected = client.post(path, files={"file": ("missing.csv", missing_phone, "text/csv")},
                           data={"dry_run": "true"})
    assert rejected.status_code == 200 and rejected.json()["errors_count"] == 1


def test_call_center_workbook_import_preserves_warranty_and_unknown_source(client):
    _seed(client)
    workbook = Workbook()
    non_warranty = workbook.active
    non_warranty.title = "غیر گارانتی"
    non_warranty.append([])
    non_warranty.append(["ردیف", "تاریخ و ساعت ثبت", "نام و نام خانوادگی", "شماره تماس", "برند خودرو", "محدوده", "مدل خودرو", "مشکل خودرو", "کنسل شد؟"])
    non_warranty.append([1, datetime(2026, 9, 1, 12), "Test A", 9123456789, "Brand", "Tehran", "X1", "Issue", True])
    non_warranty.append([None, None, None, None, None, None, None, None, False])
    warranty = workbook.create_sheet("گارانتی")
    warranty.append(["ردیف", "تاریخ و ساعت ثبت", "نام و نام خانوادگی", "شماره تماس", "برند خودرو", "مدل خودرو"])
    warranty.append([1, None, None, 9351234567, "Brand", "X2"])
    warranty.append([2, None, "No Phone", None, "Brand", "X3"])
    buffer = BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    path = "/api/v1/call-center/calls/import-workbook"
    files = lambda: {"file": ("calls.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview = client.post(path, files=files(), data={"dry_run": "true", "site_id": "demo"})
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows_valid"] == 2 and preview.json()["rows_imported"] == 0
    assert preview.json()["sheets"]["گارانتی"]["missing_phone"] == 1
    assert preview.json()["sheets"]["گارانتی"]["missing_date"] == 1
    done = client.post(path, files=files(), data={"dry_run": "false", "site_id": "demo"})
    assert done.status_code == 200 and done.json()["rows_imported"] == 2
    assert client.post(path, files=files(), data={"dry_run": "false", "site_id": "demo"}).json()["rows_skipped"] == 2
    calls = client.get("/api/v1/call-center/calls").json()["items"]
    assert {row["source"] for row in calls} == {"unknown"}
    assert {row["phone"] for row in calls} == {"09123456789", "09351234567"}
    assert {row["warranty"] for row in calls} == {True, False}
    assert {row["status"] for row in calls} == {"cancelled", "unreviewed"}
    non_warranty.cell(3, 4, 9999999999)
    changed_buffer = BytesIO()
    workbook.save(changed_buffer)
    changed = client.post(path, files={"file": ("changed.xlsx", changed_buffer.getvalue())}, data={"dry_run": "true"})
    assert changed.json()["rows_changed"] == 1
    assert changed.json()["conflicts"] == [{"sheet": "غیر گارانتی", "row": 3}]


def test_memory_and_ai_orchestrator_endpoints(client):
    _seed(client)
    m = client.get("/api/v1/sites/demo/memory").json()
    assert m["business_rules"] == [] and m["tone"] == {}
    r = client.put("/api/v1/sites/demo/memory", json={"business_rules": ["فقط تهران"], "tone": {"voice": "friendly"}})
    assert r.status_code == 200 and r.json()["business_rules"] == ["فقط تهران"]
    ctx = client.get("/api/v1/sites/demo/memory/context").json()["messages"]
    assert ctx and ctx[0]["role"] == "system" and "فقط تهران" in ctx[0]["content"]

    assert client.get("/api/v1/ai/routes").json()["providers"] == ["echo"]
    r = client.post("/api/v1/ai/sites/demo/run", json={"kind": "brief", "prompt": "write", "json_keys": ["title", "h1"],
                                                       "learn_pattern": "brief ok", "learn_evidence": "test"})
    body = r.json()
    assert r.status_code == 200 and body["ok"] and body["memory_used"] and body["response"]["parsed"] == {"title": "echo:title", "h1": "echo:h1"}
    assert client.get("/api/v1/sites/demo/memory").json()["successful_patterns"][0]["pattern"] == "brief ok"
    assert client.post("/api/v1/ai/sites/demo/run", json={"kind": "nope", "prompt": "x"}).status_code == 422


def test_jobs_endpoints(client):
    r = client.post("/api/v1/jobs", json={"type": "noop", "payload": {"site_id": "demo", "x": 1}})
    assert r.status_code == 202
    run_id = r.json()["run_id"]
    import time
    for _ in range(50):
        st = client.get(f"/api/v1/jobs/{run_id}").json()
        if st["status"] in ("succeeded", "failed"):
            break
        time.sleep(0.02)
    assert st["status"] == "succeeded" and st["result"]["echo"]["site_id"] == "demo" and st["result"]["echo"]["x"] == 1
    assert st["result"]["echo"]["job_id"] == run_id       # the queue exposes the job's own run id to every handler payload
    assert client.post("/api/v1/jobs", json={"type": "unknown"}).status_code == 422
    assert client.get("/api/v1/jobs/none").status_code == 404


def test_api_token_enforced_when_set(tmp_path, monkeypatch):
    eng = make_engine("sqlite:///" + (tmp_path / "tok.db").as_posix()); migrate(eng)
    monkeypatch.setenv("API_TOKEN", "s3cret")
    app = create_app(); app.dependency_overrides[deps.engine] = lambda: eng
    from seo_brain.ai.gateway import Gateway as _GW
    app.dependency_overrides[deps.gateway] = (lambda g: (lambda: g))(_GW(eng))
    c = TestClient(app)
    assert c.get("/api/v1/health").status_code == 200                # health is open
    assert c.get("/api/v1/sites").status_code == 401
    assert c.get("/api/v1/sites", headers={"X-API-Token": "wrong"}).status_code == 401
    assert c.get("/api/v1/sites", headers={"X-API-Token": "s3cret"}).status_code == 200


def test_legacy_dashboard_mounted(client):
    r = client.get("/legacy/api/sites")
    assert r.status_code == 200
    assert client.get("/").json()["legacy_dashboard"] == "/legacy"


def test_error_envelope_and_request_id(client):
    r = client.get("/api/v1/sites/nope", headers={"X-Request-ID": "req-1"})
    assert r.status_code == 404 and r.headers["X-Request-ID"] == "req-1"
    assert r.json() == {"error": {"code": "not_found", "message": "unknown site_id 'nope'", "details": None, "request_id": "req-1"}}
    r = client.post("/api/v1/sites", json={"site_id": "BAD", "name": "x", "canonical_url": "nope"})
    body = r.json()["error"]
    assert r.status_code == 422 and body["code"] == "validation_error" and isinstance(body["details"], list) and body["details"][0]["loc"]
    assert client.get("/api/v1/health").headers.get("X-Request-ID")


def test_delete_site_refuses_then_forces(client):
    _seed(client)
    client.put("/api/v1/sites/demo/memory", json={"tone": {"voice": "x"}})
    r = client.delete("/api/v1/sites/demo")
    assert r.status_code == 409 and r.json()["error"]["code"] == "site_has_data"
    assert r.json()["error"]["details"]["graph_nodes"] == 3 and r.json()["error"]["details"]["site_memory"] == 1
    r = client.delete("/api/v1/sites/demo?force=true")
    assert r.status_code == 200 and r.json()["deleted"] == "demo" and r.json()["related_rows_deleted"]["graph_edges"] == 2
    assert client.get("/api/v1/sites/demo").status_code == 404
    assert client.get("/api/v1/sites").json() == []
    # a site without data deletes without force
    client.post("/api/v1/sites", json={"site_id": "empty", "name": "E", "canonical_url": "https://e.example/"})
    assert client.delete("/api/v1/sites/empty").status_code == 200


def test_graph_modes_view_and_details(client):
    _seed(client)
    modes = client.get("/api/v1/sites/demo/graph/modes").json()
    assert [m["key"] for m in modes] == ["seo", "content", "links", "planner"] and all(m["title_fa"] for m in modes)
    v = client.get("/api/v1/sites/demo/graph/view", params={"mode": "seo"}).json()
    assert v["mode"]["key"] == "seo" and {n["id"] for n in v["nodes"]} == {"site:demo", "page:https://demo.example/a", "query:امداد"}
    assert {(e["source"], e["relation_type"], e["target"]) for e in v["edges"]} == {("site:demo", "HAS_PAGE", "page:https://demo.example/a"), ("page:https://demo.example/a", "RANKS_FOR", "query:امداد")}
    assert v["stats"]["by_type"] == {"SITE": 1, "PAGE": 1, "QUERY": 1} and v["truncated"] is False
    # links mode: only page-ish nodes and LINKS_TO edges (none seeded) → nodes without edges dropped when include_isolated=false
    lv = client.get("/api/v1/sites/demo/graph/view", params={"mode": "links", "include_isolated": "false"}).json()
    assert lv["nodes"] == [] and lv["edges"] == []
    lv2 = client.get("/api/v1/sites/demo/graph/view", params={"mode": "links"}).json()
    assert [n["id"] for n in lv2["nodes"]] == ["page:https://demo.example/a"]
    # types filter narrows within the mode; unknown mode → 422; limit → truncated flag
    tv = client.get("/api/v1/sites/demo/graph/view", params={"mode": "seo", "types": "QUERY"}).json()
    assert [n["type"] for n in tv["nodes"]] == ["QUERY"] and tv["edges"] == []
    assert client.get("/api/v1/sites/demo/graph/view", params={"mode": "nope"}).status_code == 422
    assert client.get("/api/v1/sites/demo/graph/view", params={"mode": "seo", "limit": 1}).json()["truncated"] is True
    # details per node kind
    d = client.get("/api/v1/sites/demo/graph/node-details/query:امداد").json()
    assert d["type"] == "QUERY" and d["keyword"]["related_pages"][0]["id"] == "page:https://demo.example/a" and d["degree"] == 1
    p = client.get("/api/v1/sites/demo/graph/node-details/page:https://demo.example/a").json()
    assert p["type"] == "PAGE" and "page" in p and p["page"]["content_status"] == "unknown" and p["related"]["queries"][0]["label"] == "امداد"
    s = client.get("/api/v1/sites/demo/graph/node-details/site:demo").json()
    assert s["type"] == "SITE" and "site" in s
    assert client.get("/api/v1/sites/demo/graph/node-details/nope:x").status_code == 404
