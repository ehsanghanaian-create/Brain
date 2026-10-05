"""Import the two-tab call-center Google Sheet export without inferring attribution.

The workbook contains no SEO/Ads column. Imported calls remain unknown until an
operator records the source. Spreadsheet row identity is stable for repeat imports;
existing calls, including operator edits, are never overwritten.
"""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import Engine, text


# Stable internal source key. Keep the private Google Sheet URL out of Git.
SHEET_ID = "call-center-sheet-v1"
SHEETS = {
    "غیر گارانتی": {"header_row": 2, "warranty": False},
    "گارانتی": {"header_row": 1, "warranty": True},
}
REQUIRED = {"تاریخ و ساعت ثبت", "نام و نام خانوادگی", "شماره تماس", "برند خودرو", "مدل خودرو"}
INSERT_COLUMNS = (
    "site_id", "occurred_at", "customer_name", "phone", "warranty", "brand", "model",
    "region", "issue", "source", "source_basis", "source_note", "campaign", "status",
    "outcome", "order_value", "follow_up_at", "source_confidence", "operator_id", "import_key",
)


def _value(raw: Any, limit: int) -> str:
    if raw is None:
        return ""
    if isinstance(raw, (float, int)) and not isinstance(raw, bool) and float(raw).is_integer():
        raw = str(int(raw))
    return str(raw).strip()[:limit]


def _phone(raw: Any) -> str:
    value = _value(raw, 40).translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
    digits = "".join(char for char in value if char.isdigit())
    # Sheets stores Iranian mobile numbers as 10-digit numeric cells when the
    # leading zero was not formatted as text. Restore only the unambiguous case.
    if len(digits) == 10 and digits.startswith("9"):
        digits = "0" + digits
    return digits[:30]


def _date(raw: Any) -> str | None:
    if isinstance(raw, datetime):
        parsed = raw
    elif isinstance(raw, str):
        try:
            parsed = datetime.fromisoformat(raw.strip().replace("Z", "+00:00"))
        except ValueError:
            return None
    else:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone(timedelta(hours=3, minutes=30)))
    return parsed.astimezone(timezone.utc).isoformat(timespec="seconds")


def import_workbook(data: bytes, eng: Engine, *, apply: bool = False,
                    site_id: str | None = None, sheet_id: str = SHEET_ID) -> dict[str, Any]:
    if len(data) > 5_000_000:
        raise ValueError("فایل XLSX باید کمتر از ۵ مگابایت باشد")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("فایل XLSX معتبر نیست") from exc
    try:
        missing = set(SHEETS) - set(workbook.sheetnames)
        if missing:
            raise ValueError("تب‌های گارانتی و غیر گارانتی در فایل پیدا نشدند")
        parsed: list[dict[str, Any]] = []
        counts: dict[str, dict[str, int]] = {}
        conflicts: list[dict[str, Any]] = []
        for sheet_name, spec in SHEETS.items():
            sheet = workbook[sheet_name]
            iterator = sheet.iter_rows(values_only=True)
            headers: list[str] = []
            for _ in range(spec["header_row"]):
                headers = [_value(value, 100) for value in next(iterator, ())]
            if not REQUIRED.issubset(headers):
                raise ValueError(f"ستون‌های ضروری تب {sheet_name} تغییر کرده‌اند")
            idx = {name: headers.index(name) for name in REQUIRED}
            for optional in ("محدوده", "مشکل خودرو", "کنسل شد؟"):
                if optional in headers:
                    idx[optional] = headers.index(optional)
            sheet_counts = counts[sheet_name] = {"candidates": 0, "valid": 0, "imported": 0,
                                                  "skipped_existing": 0, "missing_date": 0,
                                                  "short_phone": 0, "future_date": 0, "cancelled": 0,
                                                  "changed_rows": 0}
            for row_number, row in enumerate(iterator, start=spec["header_row"] + 1):
                get = lambda key: row[idx[key]] if key in idx and idx[key] < len(row) else None
                name = _value(get("نام و نام خانوادگی"), 160)
                if name == "-":
                    name = ""
                phone = _phone(get("شماره تماس"))
                if not name and not phone:
                    continue  # Formatted blank rows are not calls.
                sheet_counts["candidates"] += 1
                occurred_at = _date(get("تاریخ و ساعت ثبت"))
                if not occurred_at:
                    sheet_counts["missing_date"] += 1
                elif datetime.fromisoformat(occurred_at) > datetime.now(timezone.utc):
                    sheet_counts["future_date"] += 1
                if phone and len(phone) < 11:
                    sheet_counts["short_phone"] += 1
                cancelled = not spec["warranty"] and get("کنسل شد؟") in (True, 1, "TRUE", "true", "بله")
                if cancelled:
                    sheet_counts["cancelled"] += 1
                values = {
                    "site_id": site_id, "occurred_at": occurred_at, "customer_name": name,
                    "phone": phone, "warranty": int(spec["warranty"]),
                    "brand": _value(get("برند خودرو"), 100), "model": _value(get("مدل خودرو"), 100),
                    "region": _value(get("محدوده"), 120), "issue": _value(get("مشکل خودرو"), 2000),
                    "source": "unknown", "source_basis": "import",
                    "source_note": f"Google Sheets / {sheet_name} / row {row_number}",
                    "campaign": "", "status": "cancelled" if cancelled else "unreviewed",
                    "outcome": "pending", "order_value": None, "follow_up_at": None,
                    "source_confidence": "unknown", "operator_id": None,
                    "import_key": f"gsh:{sha256(sheet_id.encode()).hexdigest()[:12]}:{'w' if spec['warranty'] else 'n'}:{row_number}",
                }
                parsed.append({"values": values, "sheet": sheet_name, "row": row_number})
                sheet_counts["valid"] += 1
        if not parsed:
            raise ValueError("هیچ تماس قابل‌ورودی در شیت پیدا نشد")
        timestamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
        with eng.begin() as cx:
            if site_id and not cx.execute(text("SELECT 1 FROM sites WHERE site_id=:id"), {"id": site_id}).first():
                raise ValueError("شناسه سایت پیدا نشد")
            for item in parsed:
                values = item["values"]
                existing = cx.execute(text("""SELECT occurred_at,customer_name,phone,warranty,brand,model,region,issue
                    FROM call_center_calls WHERE import_key=:import_key"""), values).mappings().first()
                if existing:
                    compared = ("occurred_at", "customer_name", "phone", "warranty", "brand", "model", "region", "issue")
                    if any(existing[key] != values[key] for key in compared):
                        counts[item["sheet"]]["changed_rows"] += 1
                        if len(conflicts) < 30:
                            conflicts.append({"sheet": item["sheet"], "row": item["row"]})
                        continue
                    counts[item["sheet"]]["skipped_existing"] += 1
                    continue
                if apply:
                    cx.execute(text("INSERT INTO call_center_calls(" + ",".join(INSERT_COLUMNS) +
                                    ",created_at,updated_at) VALUES (" + ",".join(":" + c for c in INSERT_COLUMNS) +
                                    ",:timestamp,:timestamp)"), {**values, "timestamp": timestamp})
                    counts[item["sheet"]]["imported"] += 1
        return {"dry_run": not apply, "sha256": sha256(data).hexdigest(), "sheets": counts,
                "rows_valid": sum(item["valid"] for item in counts.values()),
                "rows_imported": sum(item["imported"] for item in counts.values()),
                "rows_skipped": sum(item["skipped_existing"] for item in counts.values()),
                "rows_changed": sum(item["changed_rows"] for item in counts.values()),
                "conflicts": conflicts,
                "source": "unknown", "source_reason": "در فایل ستونی برای SEO/Ads وجود ندارد"}
    finally:
        workbook.close()
