"""Import the two call-center Google Sheet tabs from an exported .xlsx file.

Default is a read-only preview. --apply writes to the local configured database.
The source file belongs under ignored data/imports, never in Git.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta, timezone
from pathlib import Path
import sys

from openpyxl import load_workbook
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from seo_brain.db.engine import get_engine  # noqa: E402
from seo_brain.db.migrate import migrate  # noqa: E402

SHEET_ID = "1GDg63z8IoNf4W6or9qAnfdpfQ0hRnx3Kn139bx_iFAY"


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def phone_text(value) -> str:
    raw = cell_text(value)
    digits = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
    phone = "".join(c for c in raw.translate(digits) if c.isdigit() or c == "+")
    if len(phone) == 10 and phone.startswith("9"):
        phone = "0" + phone
    return phone[:30]


def records(path: Path, selected: str = "all"):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = [("غیر گارانتی", False, 2), ("گارانتی", True, 1)]
    for name, warranty, header_row in sheets:
        if selected != "all" and selected != name:
            continue
        if name not in workbook:
            raise ValueError(f"missing sheet: {name}")
        sheet = workbook[name]
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_number <= header_row:
                continue
            values = list(row) + [None] * (10 - len(row))
            if warranty:
                _, raw_date, customer, phone, brand, model, *_ = values
                region, issue, cancelled = "", "", False
            else:
                _, raw_date, customer, phone, brand, region, model, issue, cancelled, *_ = values
            if not any(cell_text(v) for v in (customer, phone, brand, region, model, issue)):
                continue
            occurred_at = raw_date.replace(tzinfo=timezone(timedelta(hours=3, minutes=30))).astimezone(timezone.utc).isoformat(timespec="seconds") if isinstance(raw_date, datetime) else None
            raw_note = "" if occurred_at or raw_date is None else f"زمان خام در شیت: {cell_text(raw_date)[:80]}"
            yield {
                "site_id": None, "occurred_at": occurred_at, "customer_name": cell_text(customer)[:160],
                "phone": phone_text(phone), "warranty": int(warranty), "brand": cell_text(brand)[:100],
                "model": cell_text(model)[:100], "region": cell_text(region)[:120],
                "issue": cell_text(issue)[:2000], "source": "unknown", "source_basis": "import",
                "source_note": raw_note, "campaign": "", "status": "cancelled" if cancelled is True else "unreviewed",
                "operator_id": None, "import_key": f"gdrive:{SHEET_ID}:{name}:{row_number}",
                "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--sheet", choices=["all", "غیر گارانتی", "گارانتی"], default="all")
    parser.add_argument("--site-id", help="associate imported calls with a configured SEO Brain site")
    parser.add_argument("--apply", action="store_true", help="write the previewed rows to the configured local database")
    args = parser.parse_args()
    rows = list(records(args.xlsx, args.sheet))
    if args.site_id:
        for row in rows:
            row["site_id"] = args.site_id
    cutoff = datetime.now(timezone.utc).isoformat(timespec="seconds")
    future = sum(bool(r["occurred_at"] and r["occurred_at"] > cutoff) for r in rows)
    print(f"Rows: {len(rows)} | dated: {sum(r['occurred_at'] is not None for r in rows)} | undated: {sum(r['occurred_at'] is None for r in rows)} | future-dated: {future}")
    print("Attribution remains unknown; no sheet row is guessed to be SEO or Ads.")
    if not args.apply:
        print("Preview only. Run with --apply to import.")
        return
    eng = get_engine()
    migrate(eng)
    columns = list(rows[0]) if rows else []
    sql = text("INSERT INTO call_center_calls (" + ",".join(columns) + ") VALUES (" + ",".join(":" + c for c in columns) + ")")
    inserted = 0
    with eng.begin() as cx:
        for row in rows:
            if cx.execute(text("SELECT 1 FROM call_center_calls WHERE import_key=:import_key"), row).first():
                continue
            cx.execute(sql, row)
            inserted += 1
    print(f"Imported: {inserted} | already present: {len(rows) - inserted}")


if __name__ == "__main__":
    main()
